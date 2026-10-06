from __future__ import annotations

import argparse
import csv
import json
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


def normalizar(texto: str) -> str:
    base = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in base if not unicodedata.combining(c)).strip().lower()


def resolver_hoja(nombre_hoja: str, hojas: list[str]) -> str:
    if nombre_hoja in hojas:
        return nombre_hoja
    objetivo = normalizar(nombre_hoja)
    for hoja in hojas:
        if normalizar(hoja) == objetivo:
            return hoja
    raise ValueError(
        f"No se encontró la hoja '{nombre_hoja}'. Disponibles: {', '.join(hojas)}"
    )


def leer_registros(excel: Path, nombre_hoja: str) -> list[dict]:
    wb = load_workbook(excel, data_only=True, read_only=True)
    try:
        hoja_real = resolver_hoja(nombre_hoja, wb.sheetnames)
        ws = wb[hoja_real]
        filas = ws.iter_rows(values_only=True)
        encabezados = [str(c).strip() if c is not None else "" for c in next(filas)]

        registros = []
        for fila in filas:
            if not fila or all(v is None or str(v).strip() == "" for v in fila):
                continue
            registro = {}
            for i, encabezado in enumerate(encabezados):
                if not encabezado:
                    continue
                valor = fila[i] if i < len(fila) else None
                registro[encabezado] = valor
            registros.append(registro)
        return registros
    finally:
        wb.close()


def filtrar_por_codigo(registros: list[dict], codigo: str) -> list[dict]:
    codigo_obj = codigo.strip().upper()
    salida = []
    for registro in registros:
        campos_codigo = [
            valor
            for clave, valor in registro.items()
            if "codigo" in normalizar(clave)
        ]
        if any(str(v).strip().upper() == codigo_obj for v in campos_codigo if v is not None):
            salida.append(registro)
    return salida


def exportar_json(registros: list[dict], salida: Path | None) -> None:
    texto = json.dumps(registros, ensure_ascii=False, indent=2, default=str)
    if salida:
        salida.write_text(texto, encoding="utf-8")
        print(f"✅ JSON generado: {salida}")
    else:
        print(texto)


def exportar_csv(registros: list[dict], salida: Path) -> None:
    if not registros:
        salida.write_text("", encoding="utf-8")
        print(f"✅ CSV generado vacío: {salida}")
        return

    columnas = list(registros[0].keys())
    with salida.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columnas)
        writer.writeheader()
        writer.writerows(registros)
    print(f"✅ CSV generado: {salida}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extrae información del inventario Excel para integrarla con otras apps."
    )
    parser.add_argument(
        "--excel",
        default="inventario_hospital_v1.xlsx",
        help="Ruta al archivo Excel de inventario.",
    )
    parser.add_argument(
        "--hoja",
        default="Equipos de Cómputo",
        help="Nombre de la hoja a consultar.",
    )
    parser.add_argument(
        "--codigo",
        default=None,
        help="Filtrar por código (ej: EQC-0004).",
    )
    parser.add_argument(
        "--formato",
        choices=["json", "csv"],
        default="json",
        help="Formato de salida.",
    )
    parser.add_argument(
        "--salida",
        default=None,
        help="Archivo de salida. Si no se define en JSON, imprime en consola.",
    )

    args = parser.parse_args()
    excel = Path(args.excel)
    if not excel.exists():
        raise FileNotFoundError(f"No existe el archivo: {excel}")

    registros = leer_registros(excel, args.hoja)
    if args.codigo:
        registros = filtrar_por_codigo(registros, args.codigo)

    if args.formato == "json":
        salida = Path(args.salida) if args.salida else None
        exportar_json(registros, salida)
    else:
        salida = Path(args.salida) if args.salida else Path("inventario_exportado.csv")
        exportar_csv(registros, salida)

    print(f"📦 Registros exportados: {len(registros)}")


if __name__ == "__main__":
    main()