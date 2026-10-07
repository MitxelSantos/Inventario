# -*- coding: utf-8 -*-
"""
GENERADOR DE FORMATOS OFICIALES
===============================
Llena las plantillas oficiales del hospital con los datos del inventario (Excel):

  - Hoja de vida de equipo de cómputo  (AATPL01.F12, .docx)  -> generar_hoja_vida()
  - Reporte de mantenimiento preventivo (AATPC03.F28, .pdf)  -> generar_mantenimiento_pdf()

Las plantillas se buscan en la carpeta "plantillas/" junto a este archivo.
Dependencias: python-docx, pypdf, reportlab, openpyxl
"""
from __future__ import annotations

import io
import re
from copy import copy
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
PLANTILLAS_DIR = BASE_DIR / "plantillas"
PLANTILLA_HOJA_VIDA = PLANTILLAS_DIR / "FORMATO_HOJA_DE_VIDA_EQUIPO-DE-COMPUTO.docx"
PLANTILLA_MANTENIMIENTO = PLANTILLAS_DIR / "FORMATO_MANTENIMIENTO_PREVENTIVO_COMPUTO.pdf"

HOJA_EQUIPOS = "Equipos de Cómputo"
HOJA_MANTENIMIENTOS = "Mantenimientos"


# ----------------------------------------------------------------------------
# Lectura del Excel
# ----------------------------------------------------------------------------
def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    s = str(v).strip()
    return "" if s.lower() in ("none", "no detectado", "no tiene", "nan") else s


def _fecha(v) -> str:
    """Devuelve dd/mm/aaaa a partir de datetime o texto aaaa-mm-dd."""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    s = _txt(v)
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%d/%m/%Y")
        except ValueError:
            pass
    return s


def leer_equipo(excel_path, codigo: str) -> dict:
    """Fila del equipo como dict {encabezado: valor}. Lanza ValueError si no existe."""
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        ws = wb[HOJA_EQUIPOS]
        filas = ws.iter_rows(values_only=True)
        enc = [str(c).strip() if c is not None else "" for c in next(filas)]
        objetivo = codigo.strip().upper()
        for fila in filas:
            if fila and str(fila[1] or "").strip().upper() == objetivo:
                return {h: v for h, v in zip(enc, fila) if h}
    finally:
        wb.close()
    raise ValueError(f"No existe el equipo {codigo} en la hoja '{HOJA_EQUIPOS}'.")


def leer_mantenimientos(excel_path, codigo: str) -> list[dict]:
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        if HOJA_MANTENIMIENTOS not in wb.sheetnames:
            return []
        filas = wb[HOJA_MANTENIMIENTOS].iter_rows(values_only=True)
        enc = [str(c).strip() if c is not None else "" for c in next(filas)]
        out = []
        for fila in filas:
            if fila and str(fila[1] or "").strip().upper() == codigo.strip().upper():
                out.append({h: v for h, v in zip(enc, fila) if h})
        return out
    finally:
        wb.close()


def listar_codigos(excel_path) -> list[str]:
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        filas = wb[HOJA_EQUIPOS].iter_rows(min_row=2, values_only=True)
        return [str(f[1]).strip() for f in filas if f and f[1]]
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# Utilidades de Word
# ----------------------------------------------------------------------------
def _celdas_unicas(row):
    vistas, out = set(), []
    for c in row.cells:
        if id(c._tc) not in vistas:
            vistas.add(id(c._tc))
            out.append(c)
    return out


def _set_cell(cell, texto: str, size_pt: float = 9):
    """Escribe texto en una celda conservando el formato de la plantilla."""
    from docx.shared import Pt

    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = texto
        for r in p.runs[1:]:
            r.text = ""
    else:
        r = p.add_run(texto)
        r.font.size = Pt(size_pt)


def _marcar_opcion(cell, opcion: str):
    """Cambia ☐ por ☒ en la opción cuyo texto empieza con `opcion`."""
    if not opcion:
        return
    opcion = opcion.upper()
    for p in cell.paragraphs:
        runs = p.runs
        completo = "".join(r.text for r in runs)
        pos, acum = [], 0
        for i, r in enumerate(runs):
            for j, ch in enumerate(r.text):
                if ch == "☐":
                    pos.append((acum + j, i, j))
            acum += len(r.text)
        for k, (g, ri, rj) in enumerate(pos):
            fin = pos[k + 1][0] if k + 1 < len(pos) else len(completo)
            etiqueta = completo[g + 1:fin].strip().upper()
            if etiqueta.startswith(opcion):
                t = runs[ri].text
                runs[ri].text = t[:rj] + "☒" + t[rj + 1:]
                return


def _set_fecha_en_celda(cell, texto: str):
    _set_cell(cell, texto)


# ----------------------------------------------------------------------------
# Hoja de vida (.docx)
# ----------------------------------------------------------------------------
_TIPO_EQUIPO = {"desktop": "TORRE", "torre": "TORRE", "laptop": "PORTÁTIL",
                "portátil": "PORTÁTIL", "portatil": "PORTÁTIL",
                "all-in-one": "TODO EN UNO", "todo en uno": "TODO EN UNO"}


def _estado_hv(estado: str) -> str:
    e = estado.lower()
    if "baja" in e:
        return "DADO DE BAJA"
    if "mantenimiento" in e or "reparaci" in e:
        return "EN REP"
    if "fuera" in e:
        return "FUERA DE SERVICIO"
    if "operativo" in e:
        return "OPERATIVO"
    return ""


def _tipo_mtto_corto(tipo: str) -> str:
    t = tipo.lower()
    if "prevent" in t:
        return "MP"
    if "correct" in t:
        return "MC"
    if "actualiz" in t:
        return "ACT"
    return tipo[:3].upper()


def generar_hoja_vida(excel_path, codigo: str, carpeta_salida,
                      plantilla=PLANTILLA_HOJA_VIDA) -> Path:
    """Genera la hoja de vida del equipo usando la plantilla oficial. Devuelve la ruta."""
    import docx

    plantilla = Path(plantilla)
    if not plantilla.exists():
        raise FileNotFoundError(f"No se encontró la plantilla: {plantilla}")

    eq = leer_equipo(excel_path, codigo)
    mttos = leer_mantenimientos(excel_path, codigo)
    d = docx.Document(str(plantilla))
    T = d.tables

    def g(k):
        return _txt(eq.get(k))

    # --- 1. Identificación
    r0, r1, r2 = (_celdas_unicas(T[1].rows[i]) for i in range(3))
    _set_cell(r0[1], g("Código"))
    _marcar_opcion(r0[3], _TIPO_EQUIPO.get(g("Tipo de Equipo").lower(), ""))
    _set_fecha_en_celda(r1[3], datetime.now().strftime("%d/%m/%Y"))
    _marcar_opcion(r2[1], _estado_hv(g("Estado Operativo")))

    # --- 2. Componentes (solo CPU sale del inventario de cómputo)
    cpu = _celdas_unicas(T[2].rows[1])
    _set_cell(cpu[1], g("Marca"))
    _set_cell(cpu[2], g("Modelo"))
    _set_cell(cpu[3], g("Serial"))
    _set_cell(cpu[4], g("Procesador"))
    _set_cell(cpu[5], g("Código"))

    # --- 3. Especificaciones técnicas
    f0, f1, f2 = (_celdas_unicas(T[3].rows[i]) for i in range(3))
    _set_cell(f0[1], g("Procesador"))
    ram = g("RAM (GB)")
    _set_cell(f1[1], f"{ram} GB" if ram else "")
    tipo_disco = g("Disco 1: Tipo").upper()
    _marcar_opcion(f1[3], "M.2" if ("NVME" in tipo_disco or "M.2" in tipo_disco) else tipo_disco)
    cap = g("Disco 1: Capacidad (GB)")
    _set_cell(f2[1], f"{cap} GB" if cap else "")
    so = g("Sistema Operativo")
    arq = g("Arquitectura SO")
    _set_cell(f2[3], f"{so} ({arq})" if so and arq else so)

    # --- 4. Red y localización
    n = [_celdas_unicas(T[4].rows[i]) for i in range(5)]
    _set_cell(n[0][1], g("Nombre Equipo"))
    _set_cell(n[0][3], g("Dirección IP"))
    _set_cell(n[2][1], g("MAC Address"))
    _set_cell(n[3][1], g("Área / Servicio"))
    _set_cell(n[3][3], g("Ubicación Específica"))
    _set_cell(n[4][1], g("Responsable / Custodio"))

    # --- 5. Licencias de software
    sw = []
    if g("Sistema Operativo"):
        sw.append((g("Sistema Operativo"), g("Arquitectura SO"),
                   g("Key Windows") or g("Licencia Windows")))
    if g("Versión Office"):
        sw.append(("Microsoft Office", g("Versión Office"), g("Licencia Office")))
    if g("Antivirus Instalado"):
        sw.append((g("Antivirus Instalado"), g("Estado Antivirus"), ""))
    if g("Descripción Software"):
        sw.append((g("Descripción Software"), "", ""))
    for i, fila in enumerate(sw[:5], start=1):
        celdas = _celdas_unicas(T[5].rows[i])
        for c, valor in zip(celdas, fila):
            _set_cell(c, valor)

    # --- 7. Historial de mantenimientos (hasta 19 filas)
    for i, m in enumerate(mttos[:19], start=1):
        celdas = _celdas_unicas(T[7].rows[i])
        _set_cell(celdas[1], _fecha(m.get("Fecha Mantenimiento")))
        _set_cell(celdas[2], _tipo_mtto_corto(_txt(m.get("Tipo Mantenimiento"))))
        _set_cell(celdas[3], _txt(m.get("Descripción Actividades")))
        _set_cell(celdas[4], _txt(m.get("Técnico Responsable")))
        _set_cell(celdas[5], _txt(m.get("N° Consecutivo")))

    carpeta = Path(carpeta_salida)
    carpeta.mkdir(parents=True, exist_ok=True)
    salida = carpeta / f"HV_{g('Código')}.docx"
    d.save(str(salida))
    return salida


# ----------------------------------------------------------------------------
# Reporte de mantenimiento preventivo (PDF, escritura sobre la plantilla)
# ----------------------------------------------------------------------------
# Coordenadas medidas sobre la plantilla (pt, origen arriba-izquierda; página 612x792).
_CHECK_TOPS = [443, 479, 504, 529, 565, 590, 626, 662, 687, 712]   # filas 1..10
_CHECK_X = {"cumple": 409, "no_cumple": 474, "na": 539}
_TIPO_X = {"TODO EN UNO": 177, "TORRE": 271, "PORTÁTIL": 331, "OTRO": 406}


def _wrap(c, texto, ancho, fuente, tam):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    lineas, actual = [], ""
    for palabra in str(texto).split():
        prueba = f"{actual} {palabra}".strip()
        if stringWidth(prueba, fuente, tam) <= ancho:
            actual = prueba
        else:
            if actual:
                lineas.append(actual)
            actual = palabra
    if actual:
        lineas.append(actual)
    return lineas


def generar_mantenimiento_pdf(excel_path, codigo: str, carpeta_salida, *,
                              numero_orden="", fecha_programada=None,
                              fecha_ejecucion=None, hora_inicio="", hora_fin="",
                              checklist=None, materiales=None,
                              observaciones="", seguimiento="",
                              plantilla=PLANTILLA_MANTENIMIENTO) -> Path:
    """
    Escribe sobre la plantilla oficial el reporte de mantenimiento preventivo.

    checklist: lista de 10 valores "cumple" | "no_cumple" | "na" | None (sin marcar)
    materiales: lista de tuplas (material, descripcion, cantidad), máx. 8
    fechas: str dd/mm/aaaa, date o datetime. Horas: "HH:MM".
    """
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas

    plantilla = Path(plantilla)
    if not plantilla.exists():
        raise FileNotFoundError(f"No se encontró la plantilla: {plantilla}")

    eq = leer_equipo(excel_path, codigo)
    H = 792.0
    fuente, negrita = "Helvetica", "Helvetica-Bold"

    def y(top_bottom):  # convierte coordenada superior a la de reportlab
        return H - top_bottom

    def f(v):
        return _fecha(v) if v else ""

    def partes_fecha(v):
        s = f(v)
        return s.split("/") if s.count("/") == 2 else ["", "", ""]

    buf1, buf2 = io.BytesIO(), io.BytesIO()
    c = canvas.Canvas(buf1, pagesize=(612, 792))
    c.setFont(fuente, 9)

    def texto(x, bottom, s, tam=9, fnt=fuente):
        c.setFont(fnt, tam)
        c.drawString(x, y(bottom) + 2, str(s))

    from reportlab.pdfbase.pdfmetrics import stringWidth

    def ajustado(x, bottom, s, ancho=126, tam=9):
        s = str(s)
        while tam > 6 and stringWidth(s, fuente, tam) > ancho:
            tam -= 0.5
        texto(x, bottom, s, tam)

    # 1. Información de la orden (celdas de valor: x=176 izquierda, x=446 derecha)
    ajustado(176, 163, numero_orden)
    d, m, a = partes_fecha(fecha_programada)
    if d:
        texto(455, 163, d)
        texto(490, 163, m)
        texto(524, 163, a)
    d, m, a = partes_fecha(fecha_ejecucion)
    if d:
        texto(184, 186, d)
        texto(218, 186, m)
        texto(250, 186, a)
    hi = (hora_inicio or "").split(":")
    if len(hi) == 2:
        texto(466, 186, hi[0])
        texto(521, 186, hi[1])
    hf = (hora_fin or "").split(":")
    if len(hf) == 2:
        texto(197, 210, hf[0])
        texto(252, 210, hf[1])
    try:
        t1 = datetime.strptime(hora_inicio, "%H:%M")
        t2 = datetime.strptime(hora_fin, "%H:%M")
        minutos = int((t2 - t1).total_seconds() // 60)
        if minutos >= 0:
            texto(492, 210, minutos)
    except (ValueError, TypeError):
        pass

    # 2. Datos del equipo
    tipo = _TIPO_EQUIPO.get(_txt(eq.get("Tipo de Equipo")).lower(), "OTRO")
    c.setFont(negrita, 10)
    c.drawString(_TIPO_X[tipo] + 1.5, y(266) + 1, "X")
    if tipo == "OTRO":
        texto(460, 267, _txt(eq.get("Tipo de Equipo")), 8)
    ajustado(176, 291, _txt(eq.get("Marca")))
    ajustado(446, 291, _txt(eq.get("Modelo")))
    ajustado(176, 314, _txt(eq.get("Serial")))
    ajustado(446, 314, _txt(eq.get("Código")))
    ajustado(176, 338, _txt(eq.get("Ubicación Específica")))
    ajustado(446, 338, _txt(eq.get("Responsable / Custodio")))
    ajustado(176, 361, _txt(eq.get("Dirección IP")))
    ajustado(446, 361, _txt(eq.get("Nombre Equipo")))

    # 3. Checklist
    for top, valor in zip(_CHECK_TOPS, checklist or []):
        x = _CHECK_X.get(valor)
        if x:
            c.setFont(negrita, 10)
            c.drawString(x + 1, y(top + 11) + 1, "X")
    c.showPage()

    # Página 2
    c.setFont(fuente, 9)
    fila_top = 166
    for mat, desc, cant in (materiales or [])[:8]:
        c.setFont(fuente, 8)
        c.drawString(44, y(fila_top + 10), str(mat)[:38])
        c.drawString(178, y(fila_top + 10), str(desc)[:62])
        c.drawString(480, y(fila_top + 10), str(cant)[:12])
        fila_top += 17
    for bloque, top, limite in ((observaciones, 328, 4), (seguimiento, 450, 6)):
        for i, linea in enumerate(_wrap(c, bloque, 515, fuente, 9)[:limite]):
            c.setFont(fuente, 9)
            c.drawString(42, y(top + 10 + i * 12), linea)
    c.showPage()
    c.save()

    # Fusionar overlay con la plantilla
    base = PdfReader(str(plantilla))
    capa = PdfReader(io.BytesIO(buf1.getvalue()))
    w = PdfWriter()
    for i, pagina in enumerate(base.pages):
        pagina.merge_page(capa.pages[i])
        w.add_page(pagina)
    carpeta = Path(carpeta_salida)
    carpeta.mkdir(parents=True, exist_ok=True)
    nombre_orden = re.sub(r"[^\w-]", "", str(numero_orden)) or datetime.now().strftime("%Y%m%d_%H%M")
    salida = carpeta / f"MP_{_txt(eq.get('Código'))}_{nombre_orden}.pdf"
    with open(salida, "wb") as fh:
        w.write(fh)
    return salida


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Genera hoja de vida (.docx) o reporte de mantenimiento (.pdf)")
    ap.add_argument("tipo", choices=["hv", "mtto"])
    ap.add_argument("codigo")
    ap.add_argument("--excel", default="inventario_hospital_v1.xlsx")
    ap.add_argument("--salida", default="salida_formatos")
    a = ap.parse_args()
    if a.tipo == "hv":
        print(generar_hoja_vida(a.excel, a.codigo, a.salida))
    else:
        hoy = datetime.now().strftime("%d/%m/%Y")
        print(generar_mantenimiento_pdf(a.excel, a.codigo, a.salida, fecha_ejecucion=hoy,
                                        checklist=["cumple"] * 10))
